from __future__ import annotations

from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook
from zipfile import BadZipFile

from persona.adapters.base import FileAdapter
from persona.core.text_ops import apply_text_replacements, strict_restore_text
from persona.exceptions import DocumentProcessingError, InputFileError
from persona.models.entities import MapEntry, RestoreStats, TextReplacement, TextSegment


class XlsxAdapter(FileAdapter):
    supported_suffixes = (".xlsx",)

    def extract_segments(self, input_path: Path) -> list[TextSegment]:
        workbook = self._open_workbook(input_path)
        segments: list[TextSegment] = []
        try:
            for worksheet in workbook.worksheets:
                for row in worksheet.iter_rows():
                    for cell in row:
                        if isinstance(cell.value, str) and not cell.value.startswith("="):
                            location = f"{worksheet.title}!{cell.coordinate}"
                            segments.append(
                                TextSegment(
                                    segment_id=location,
                                    location=location,
                                    text=cell.value,
                                    container_type="cell",
                                )
                            )
        except Exception as exc:
            raise DocumentProcessingError("Unable to extract text cells from XLSX input.") from exc
        return segments

    def apply_replacements(
        self,
        input_path: Path,
        output_path: Path,
        replacements_by_segment: dict[str, list[TextReplacement]],
    ) -> None:
        workbook = self._open_workbook(input_path)
        try:
            for worksheet in workbook.worksheets:
                for row in worksheet.iter_rows():
                    for cell in row:
                        location = f"{worksheet.title}!{cell.coordinate}"
                        replacements = replacements_by_segment.get(location)
                        if replacements and isinstance(cell.value, str):
                            cell.value = apply_text_replacements(cell.value, replacements)
            output_path.parent.mkdir(parents=True, exist_ok=True)
            workbook.save(str(output_path))
        except Exception as exc:
            raise DocumentProcessingError("Unable to write anonymized XLSX output.") from exc

    def restore_file(
        self,
        censored_path: Path,
        output_path: Path,
        entries: list[MapEntry],
        *,
        root_key: bytes | None = None,
    ) -> RestoreStats:
        workbook = self._open_workbook(censored_path)
        entries_by_segment: dict[str, list[MapEntry]] = defaultdict(list)
        for entry in entries:
            entries_by_segment[entry.segment_id].append(entry)

        stats = RestoreStats()
        try:
            for worksheet in workbook.worksheets:
                for row in worksheet.iter_rows():
                    for cell in row:
                        location = f"{worksheet.title}!{cell.coordinate}"
                        segment_entries = entries_by_segment.get(location)
                        if not segment_entries or not isinstance(cell.value, str):
                            continue
                        outcome = strict_restore_text(
                            cell.value,
                            segment_entries,
                            root_key=root_key,
                            segment_id=location,
                            location=location,
                        )
                        if outcome.text != cell.value:
                            cell.value = outcome.text
                        stats.extend(outcome.stats)
            output_path.parent.mkdir(parents=True, exist_ok=True)
            workbook.save(str(output_path))
        except Exception as exc:
            raise DocumentProcessingError("Unable to restore XLSX output.") from exc
        return stats

    def _open_workbook(self, input_path: Path):
        try:
            return load_workbook(filename=str(input_path))
        except FileNotFoundError as exc:
            raise InputFileError(f"XLSX file not found: {input_path}") from exc
        except BadZipFile as exc:
            raise InputFileError(f"XLSX file is not a valid Office workbook: {input_path}") from exc
        except Exception as exc:
            raise DocumentProcessingError(f"Unable to open XLSX file: {input_path}") from exc
