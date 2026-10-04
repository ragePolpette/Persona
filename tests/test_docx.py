from __future__ import annotations

import zipfile
from pathlib import Path

import pytest
from docx import Document as WordDocument
from typer.testing import CliRunner

from persona import cli
from persona.documents import DocxDocument, open_document
from persona.engine import analyze, apply, verify
from persona.exceptions import InputError
from persona.restore import placeholders_in, restore_segments
from persona.vault import Argon2Params, Vault
from tests.conftest import FAST_KDF
from tests.docx_fixtures import SENSITIVE, all_xml_text, build_docx


@pytest.fixture
def source(tmp_path: Path) -> Path:
    return build_docx(tmp_path / "contratto.docx")


def anonymize(source: Path, vault: Vault, out: Path) -> Path:
    document = open_document(source)
    result = apply(analyze(document.segments, vault), vault)
    document.write(out, result.edits)
    return out


def test_segments_reach_every_hiding_place(source: Path) -> None:
    ids = {segment.id: segment.text for segment in open_document(source).segments}
    texts = "\n".join(ids.values())
    assert any("header" in i and "Mario Rossi" in t for i, t in ids.items())
    assert any("footer" in i and "mario.rossi@example.test" in t for i, t in ids.items())
    assert any("footnotes" in i and "Giulia Marchetti" in t for i, t in ids.items())
    assert any("comments" in i and "Verificare con Mario Rossi" in t for i, t in ids.items())
    assert "Nel riquadro: Mario Rossi" in texts  # text box
    assert "vecchio referente Mario Rossi" in texts  # tracked deletion
    assert 'HYPERLINK "mailto:mario.rossi@example.test"' in texts  # field code
    assert any(i.endswith("rId" + i.rsplit("rId", 1)[-1]) and "linkedin.com" in t for i, t in ids.items())  # link target
    assert "Contratto Tessitura Valdarno S.r.l." in texts  # document title
    assert "Contratto con Mario Rossi, telefono 338 4567890." in texts  # text split over runs
    assert "Giulia Marchetti" in texts and "Rossi" in texts


def test_nothing_sensitive_survives_anywhere_in_the_package(source: Path, vault: Vault, tmp_path: Path) -> None:
    assert all(item in all_xml_text(source) for item in SENSITIVE if item != "linkedin.com/in/mario-rossi")
    out = anonymize(source, vault, tmp_path / "out.docx")
    package = all_xml_text(out)
    for item in SENSITIVE:
        assert item not in package, f"leaked: {item!r}"
    for fragment in ("Mario", "Rossi", "Marchetti", "Valdarno"):
        assert fragment not in package, f"leaked fragment: {fragment!r}"
    assert verify({s.id: s.text for s in open_document(out).segments}, vault) == []


def test_identity_metadata_thumbnail_and_authors_are_scrubbed(source: Path, vault: Vault, tmp_path: Path) -> None:
    out = anonymize(source, vault, tmp_path / "out.docx")
    with zipfile.ZipFile(out) as archive:
        names = archive.namelist()
        assert "docProps/thumbnail.jpeg" not in names
        assert b"thumbnail" not in archive.read("_rels/.rels")
        assert b"Mario" not in archive.read("word/people.xml")
        assert b'w:author="Author"' in archive.read("word/comments.xml")
        assert b'w:author="Author"' in archive.read("word/document.xml")
    reopened = WordDocument(str(out))  # still a valid Word file
    assert reopened.core_properties.author == ""
    assert reopened.core_properties.last_modified_by == ""


def test_business_content_is_untouched(source: Path, vault: Vault, tmp_path: Path) -> None:
    out = anonymize(source, vault, tmp_path / "out.docx")
    assert "Valori business: 12.500,00 euro entro 30 giorni." in [p.text for p in WordDocument(str(out)).paragraphs]


def test_formatting_survives_when_a_name_spans_runs(source: Path, vault: Vault, tmp_path: Path) -> None:
    out = anonymize(source, vault, tmp_path / "out.docx")
    paragraph = next(p for p in WordDocument(str(out)).paragraphs if p.text.startswith("Contratto con"))
    assert "[PERSONA_" in paragraph.text and paragraph.text.endswith(", telefono [TELEFONO_1].")
    runs = {run.text: run.bold for run in paragraph.runs if run.text}
    assert runs["Contratto con "] is None
    assert any(bold for text, bold in runs.items() if "[PERSONA_" in text)  # placeholder kept the bold run


def test_restore_gives_back_the_original_text(source: Path, vault: Vault, tmp_path: Path) -> None:
    out = anonymize(source, vault, tmp_path / "out.docx")
    censored = open_document(out)
    edits, report = restore_segments(censored.segments, vault, expected=set().union(*(placeholders_in(s.text) for s in censored.segments)))
    restored_path = tmp_path / "restored.docx"
    censored.write(restored_path, edits)
    assert report.clean
    original = {s.id: s.text for s in open_document(source).segments if "docProps" not in s.id}
    restored = {s.id: s.text for s in open_document(restored_path).segments if "docProps" not in s.id}
    assert restored == original


def test_restore_finds_a_placeholder_the_ai_split_over_runs(vault: Vault, tmp_path: Path) -> None:
    vault.placeholder_for("PERSONA", "Mario Rossi")
    ai = WordDocument()
    paragraph = ai.add_paragraph("Gentile ")
    paragraph.add_run("[PERS").bold = True
    paragraph.add_run("ONA_1]").italic = True
    paragraph.add_run(", grazie.")
    path = tmp_path / "ai.docx"
    ai.save(str(path))
    document = open_document(path)
    edits, report = restore_segments(document.segments, vault)
    out = tmp_path / "ai.restored.docx"
    document.write(out, edits)
    assert [p.text for p in WordDocument(str(out)).paragraphs][0] == "Gentile Mario Rossi, grazie."
    assert report.clean


def test_images_and_other_unreadable_content_produce_a_warning(tmp_path: Path) -> None:
    path = build_docx(tmp_path / "with_image.docx", with_image=True)
    warnings = open_document(path).warnings
    assert any("image" in w and "NOT anonymized" in w for w in warnings)
    assert open_document(build_docx(tmp_path / "plain.docx")).warnings == []


def test_invalid_files_are_clean_errors(tmp_path: Path) -> None:
    bad = tmp_path / "bad.docx"
    bad.write_bytes(b"not a zip")
    with pytest.raises(InputError, match="not a valid DOCX"):
        open_document(bad)
    empty_zip = tmp_path / "empty.docx"
    with zipfile.ZipFile(empty_zip, "w") as archive:
        archive.writestr("hello.txt", "x")
    with pytest.raises(InputError, match="no word/document.xml"):
        DocxDocument(empty_zip)


def test_docx_through_the_cli(source: Path, tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setenv("PERSONA_HOME", str(tmp_path / "home"))
    monkeypatch.setenv("PERSONA_PASSWORD", "pw")
    monkeypatch.setattr(Argon2Params.__init__, "__defaults__", (1, 8, 1, 32))
    runner = CliRunner()
    assert runner.invoke(cli.app, ["init", "-p", "d"]).exit_code == 0
    done = runner.invoke(cli.app, ["anonymize", str(source), "-p", "d"])
    assert done.exit_code == 0, done.output
    anon = tmp_path / "contratto.anon.docx"
    assert "Mario" not in all_xml_text(anon)
    assert not list(tmp_path.glob(".*staging*"))
    assert runner.invoke(cli.app, ["verify", str(anon), "-p", "d"]).exit_code == 0
    leaky = runner.invoke(cli.app, ["verify", str(source), "-p", "d"])
    assert leaky.exit_code == 2

    back = runner.invoke(cli.app, ["restore", str(anon), "--sent", str(anon), "-p", "d"])
    assert back.exit_code == 0, back.output
    restored = WordDocument(str(tmp_path / "contratto.anon.restored.docx"))
    assert "Fornitore: Tessitura Valdarno S.r.l., referente Dott.ssa Giulia Marchetti." in [p.text for p in restored.paragraphs]


def test_custom_xml_text_is_masked_too(tmp_path: Path, vault: Vault) -> None:
    path = build_docx(tmp_path / "c.docx")
    with zipfile.ZipFile(path) as archive:
        entries = {n: archive.read(n) for n in archive.namelist()}
    entries["customXml/item9.xml"] = b'<?xml version="1.0"?><meta><owner>Mario Rossi</owner><id>12</id></meta>'
    with zipfile.ZipFile(path, "w") as archive:
        for name, data in entries.items():
            archive.writestr(name, data)
    out = anonymize(path, vault, tmp_path / "out.docx")
    assert "Mario" not in all_xml_text(out)
    assert b"<id>12</id>" in zipfile.ZipFile(out).read("customXml/item9.xml")
