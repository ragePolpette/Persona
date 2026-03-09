from __future__ import annotations

import base64
import json
import os
from pathlib import Path

import pytest
import spacy
from cryptography.hazmat.primitives.ciphers.aead import AESGCM
from docx import Document
from typer.testing import CliRunner

from persona.cli import app
from persona.core.pipeline import anonymize_file
from persona.core.placeholders import build_placeholder, compute_placeholder_integrity_tag
from persona.core.text_ops import strict_restore_text
from persona.exceptions import (
    CorruptKeystoreError,
    CorruptMapError,
    InputFileError,
    KeystoreDecryptionError,
    MapDecryptionError,
)
from persona.models.entities import DocumentMetadata, MapEntry
from persona.security.argon2_utils import Argon2Params, derive_key
from persona.security.keystore import create_keystore, load_root_key
from persona.security.map_store import AAD as MAP_AAD, decrypt_map_file, encrypt_map_file


RUNNER = CliRunner()
FAST_PARAMS = Argon2Params(time_cost=1, memory_cost_kib=8192, parallelism=1, hash_len=32)
ROOT_KEY = b"0123456789abcdef0123456789abcdef"


def _build_test_model(tmp_path: Path) -> Path:
    model_path = tmp_path / "test_spacy_model"
    nlp = spacy.blank("en")
    ruler = nlp.add_pipe("entity_ruler")
    ruler.add_patterns([{"label": "PERSON", "pattern": "Mario Rossi"}])
    nlp.to_disk(model_path)
    return model_path


def _build_placeholder(token_id: str, masked_value: str, root_key: bytes = ROOT_KEY) -> str:
    return build_placeholder(
        token_id,
        masked_value,
        integrity_tag=compute_placeholder_integrity_tag(root_key, token_id, masked_value),
    )


def _sample_entry(placeholder: str) -> MapEntry:
    return MapEntry(
        token_id="ABCDEF123456",
        entity_type="PERSON",
        original_value="Mario Rossi",
        masked_value="Tario Xossi",
        placeholder=placeholder,
        location="body/p0",
        segment_id="body/p0",
        start=0,
        end=11,
    )


def _write_corrupted_ciphertext(path: Path, payload: dict[str, object], field: str) -> None:
    mutated = dict(payload)
    raw = bytearray(base64.b64decode(str(mutated[field]).encode("ascii")))
    raw[0] ^= 0x01
    mutated[field] = base64.b64encode(bytes(raw)).decode("ascii")
    path.write_text(json.dumps(mutated), encoding="utf-8")


def test_wrong_password_on_keystore_raises_domain_error(tmp_path: Path) -> None:
    keystore_path = tmp_path / "keystore.json"
    create_keystore("correct-password", path=keystore_path, params=FAST_PARAMS)

    with pytest.raises(KeystoreDecryptionError):
        load_root_key("wrong-password", path=keystore_path)


def test_malformed_keystore_json_raises_domain_error(tmp_path: Path) -> None:
    keystore_path = tmp_path / "keystore.json"
    keystore_path.write_text("{not-json", encoding="utf-8")

    with pytest.raises(CorruptKeystoreError):
        load_root_key("password", path=keystore_path)


def test_altered_keystore_ciphertext_raises_domain_error(tmp_path: Path) -> None:
    keystore_path = tmp_path / "keystore.json"
    create_keystore("correct-password", path=keystore_path, params=FAST_PARAMS)
    payload = json.loads(keystore_path.read_text(encoding="utf-8"))
    _write_corrupted_ciphertext(keystore_path, payload, "ciphertext_b64")

    with pytest.raises(KeystoreDecryptionError):
        load_root_key("correct-password", path=keystore_path)


def test_wrong_password_on_map_raises_domain_error(tmp_path: Path) -> None:
    map_path = tmp_path / "sample.persona-map.json"
    document = DocumentMetadata("sample.docx", "C:/tmp/sample.docx", ".docx", "abc")
    encrypt_map_file(map_path, "correct-password", document, [], params=FAST_PARAMS)

    with pytest.raises(MapDecryptionError):
        decrypt_map_file(map_path, "wrong-password")


def test_malformed_map_json_raises_domain_error(tmp_path: Path) -> None:
    map_path = tmp_path / "sample.persona-map.json"
    map_path.write_text("{not-json", encoding="utf-8")

    with pytest.raises(CorruptMapError):
        decrypt_map_file(map_path, "password")


def test_altered_map_ciphertext_raises_domain_error(tmp_path: Path) -> None:
    map_path = tmp_path / "sample.persona-map.json"
    document = DocumentMetadata("sample.docx", "C:/tmp/sample.docx", ".docx", "abc")
    encrypt_map_file(map_path, "correct-password", document, [], params=FAST_PARAMS)
    payload = json.loads(map_path.read_text(encoding="utf-8"))
    _write_corrupted_ciphertext(map_path, payload, "ciphertext_b64")

    with pytest.raises(MapDecryptionError):
        decrypt_map_file(map_path, "correct-password")


def test_malformed_decrypted_map_payload_raises_domain_error(tmp_path: Path) -> None:
    map_path = tmp_path / "sample.persona-map.json"
    salt = os.urandom(16)
    nonce = os.urandom(12)
    key = derive_key("password", salt, params=FAST_PARAMS)
    malformed_payload = json.dumps(
        {
            "version": 1,
            "document": {
                "source_name": "sample.docx",
                "source_path": "C:/tmp/sample.docx",
                "file_format": ".docx",
                "content_sha256": "abc",
            }
        }
    ).encode("utf-8")
    ciphertext = AESGCM(key).encrypt(nonce, malformed_payload, MAP_AAD)
    map_path.write_text(
        json.dumps(
            {
                "version": 1,
                "kdf": FAST_PARAMS.to_dict() | {"salt_b64": base64.b64encode(salt).decode("ascii")},
                "nonce_b64": base64.b64encode(nonce).decode("ascii"),
                "ciphertext_b64": base64.b64encode(ciphertext).decode("ascii"),
            }
        ),
        encoding="utf-8",
    )

    with pytest.raises(CorruptMapError):
        decrypt_map_file(map_path, "password")


def test_placeholder_malformed_is_reported() -> None:
    entry = _sample_entry(_build_placeholder("ABCDEF123456", "Tario Xossi"))
    outcome = strict_restore_text(
        "broken [[P2|ABCDEF123456|BAD",
        [entry],
        root_key=ROOT_KEY,
        segment_id="body/p0",
        location="body/p0",
    )

    assert outcome.stats.untouched_invalid_count == 1
    assert any(issue.code == "MALFORMED_PLACEHOLDER" for issue in outcome.stats.issues)


def test_placeholder_parseable_but_unexpected_is_left_untouched() -> None:
    expected = _sample_entry(_build_placeholder("ABCDEF123456", "Tario Xossi"))
    unexpected_placeholder = _build_placeholder("ZZZZZZ123456", "Other Person")
    outcome = strict_restore_text(
        f"prefix {unexpected_placeholder}",
        [expected],
        root_key=ROOT_KEY,
        segment_id="body/p0",
        location="body/p0",
    )

    assert outcome.text.endswith(unexpected_placeholder)
    assert outcome.stats.untouched_invalid_count == 1
    assert any(issue.code == "UNEXPECTED_PLACEHOLDER" for issue in outcome.stats.issues)


def test_missing_expected_placeholder_is_reported() -> None:
    entry = _sample_entry(_build_placeholder("ABCDEF123456", "Tario Xossi"))
    outcome = strict_restore_text(
        "no placeholders here",
        [entry],
        root_key=ROOT_KEY,
        segment_id="body/p0",
        location="body/p0",
    )

    assert outcome.stats.restored_count == 0
    assert outcome.stats.missing_expected_count == 1
    assert any(issue.code == "MISSING_EXPECTED_PLACEHOLDER" for issue in outcome.stats.issues)


def test_duplicate_placeholder_in_same_segment_is_not_fully_restored() -> None:
    placeholder = _build_placeholder("ABCDEF123456", "Tario Xossi")
    entry = _sample_entry(placeholder)
    outcome = strict_restore_text(
        f"{placeholder} {placeholder}",
        [entry],
        root_key=ROOT_KEY,
        segment_id="body/p0",
        location="body/p0",
    )

    assert outcome.stats.restored_count == 1
    assert outcome.stats.untouched_invalid_count == 1
    assert any(issue.code == "DUPLICATE_PLACEHOLDER" for issue in outcome.stats.issues)


def test_review_reprompts_after_invalid_manual_payload() -> None:
    from persona.models.entities import DetectionMatch
    from persona.review.interactive import review_matches
    from rich.console import Console

    matches = [
        DetectionMatch(
            match_id="m0001",
            segment_id="body/p0",
            location="body/p0",
            entity_type="PERSON",
            original_value="Mario Rossi",
            start=0,
            end=11,
            score=0.9,
            context="[Mario Rossi]",
            token_id="ABCDEF123456",
            masked_value="Tario Xossi",
            placeholder=_build_placeholder("ABCDEF123456", "Tario Xossi"),
        )
    ]
    answers = iter(["e", "bad|value", "e", "Valid Mask"])

    reviewed = review_matches(matches, prompt=lambda _: next(answers), console=Console(record=True))

    assert reviewed[0].masked_value == "Valid Mask"


def test_supported_suffix_with_invalid_content_fails_cleanly(tmp_path: Path) -> None:
    invalid_docx = tmp_path / "invalid.docx"
    invalid_docx.write_text("not a zip package", encoding="utf-8")

    with pytest.raises(InputFileError):
        anonymize_file(invalid_docx, password="secret", review=False, enabled_entities=("PERSON",))


def test_cli_returns_decryption_exit_code_on_wrong_map_password(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    input_path = tmp_path / "source.docx"
    model_path = _build_test_model(tmp_path)
    monkeypatch.setenv("PERSONA_KEYSTORE_PATH", str(tmp_path / "keystore.json"))
    document = Document()
    document.add_paragraph("Mario Rossi")
    document.save(input_path)
    anonymize_result = anonymize_file(
        input_path=input_path,
        password="correct-password",
        out_dir=tmp_path,
        review=False,
        enabled_entities=("PERSON",),
        spacy_model=str(model_path),
    )
    report_path = tmp_path / "restore-report.json"

    result = RUNNER.invoke(
        app,
        [
            "restore",
            str(anonymize_result.output_file),
            "--map",
            str(anonymize_result.map_file),
            "--password-env",
            "PERSONA_PASSWORD",
            "--report-json",
            str(report_path),
        ],
        env={"PERSONA_PASSWORD": "wrong-password", "PERSONA_KEYSTORE_PATH": str(tmp_path / "keystore.json")},
    )

    assert result.exit_code == 3
    assert "Unable to decrypt encrypted map" in result.stdout
    assert not report_path.exists()


def test_cli_returns_restore_integrity_exit_code_on_partial_restore(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    input_path = tmp_path / "source.docx"
    model_path = _build_test_model(tmp_path)
    monkeypatch.setenv("PERSONA_KEYSTORE_PATH", str(tmp_path / "keystore.json"))
    document = Document()
    document.add_paragraph("Mario Rossi")
    document.save(input_path)
    anonymize_result = anonymize_file(
        input_path=input_path,
        password="correct-password",
        out_dir=tmp_path,
        review=False,
        enabled_entities=("PERSON",),
        spacy_model=str(model_path),
    )
    tampered = Document(anonymize_result.output_file)
    tampered.paragraphs[0].text = tampered.paragraphs[0].text.replace("|", "!", 1)
    tampered.save(anonymize_result.output_file)

    result = RUNNER.invoke(
        app,
        [
            "restore",
            str(anonymize_result.output_file),
            "--map",
            str(anonymize_result.map_file),
            "--password-env",
            "PERSONA_PASSWORD",
        ],
        env={"PERSONA_PASSWORD": "correct-password", "PERSONA_KEYSTORE_PATH": str(tmp_path / "keystore.json")},
    )

    assert result.exit_code == 4
    assert "Restored file:" in result.stdout
    assert "integrity issues" in result.stdout


def test_cli_returns_password_resolution_error_when_env_missing() -> None:
    result = RUNNER.invoke(
        app,
        [
            "restore",
            "missing.docx",
            "--map",
            "missing.json",
            "--password-env",
            "PERSONA_PASSWORD",
        ],
        env={},
    )

    assert result.exit_code == 3
    assert "Environment variable 'PERSONA_PASSWORD' is not set." in result.stdout
