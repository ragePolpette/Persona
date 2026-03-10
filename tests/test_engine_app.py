from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document
from fastapi.testclient import TestClient

from persona.engine.analysis import BlockAnalysisOptions, analyze_segments_with_backend
from persona.engine.llm import LLMBackendConfig, LLMOutputValidationError, MockLLMBackend, parse_backend_response
from persona.engine.service import PersonaEngine
from persona.engine.storage import LocalDocumentStore, StoredDocument
from persona.models.entities import TextSegment
from persona.web.app import create_web_app


def _build_docx(path: Path, text: str) -> None:
    document = Document()
    document.add_paragraph(text)
    document.save(path)


def test_parse_backend_response_accepts_multi_block_json() -> None:
    text = "Alice Example called Bob."
    raw = """
    {
      "findings": [
        {"start": 0, "end": 13, "text": "Alice Example", "confidence": 0.91, "label": "person"},
        {"start": 21, "end": 24, "text": "Bob", "confidence": 0.77}
      ]
    }
    """

    findings = parse_backend_response(raw, text)

    assert len(findings) == 2
    assert findings[0].text == "Alice Example"
    assert findings[1].text == "Bob"


def test_parse_backend_response_rejects_invalid_offsets() -> None:
    with pytest.raises(LLMOutputValidationError):
        parse_backend_response('{"findings":[{"start": 0, "end": 40, "text":"oops"}]}', "short text")


def test_block_analysis_prefers_larger_overlapping_block() -> None:
    segment = TextSegment(
        segment_id="body/p0",
        location="body/p0",
        container_type="paragraph",
        text="Alice Example met Charlie.",
    )
    backend = MockLLMBackend(
        static_findings=[
            # chunk-local spans
            # "Alice"
            type("Tmp", (), {"start": 0, "end": 5, "text": "Alice", "confidence": 0.60, "label": "person", "reason": None})(),
            type("Tmp", (), {"start": 0, "end": 13, "text": "Alice Example", "confidence": 0.92, "label": "person", "reason": "full name"})(),
        ]
    )

    matches = analyze_segments_with_backend([segment], backend, options=BlockAnalysisOptions(max_chunk_chars=256))

    assert len(matches) == 1
    assert matches[0].original_value == "Alice Example"
    assert matches[0].reason == "full name"


def test_engine_end_to_end_uses_block_detection_and_restore(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    input_path = tmp_path / "source.docx"
    _build_docx(input_path, "Alice Example works at Secret Corp.")
    monkeypatch.setenv("PERSONA_KEYSTORE_PATH", str(tmp_path / "keystore.json"))

    def response_factory(text: str):
        findings = []
        for needle, label in (("Alice Example", "person"), ("Secret Corp", "organization")):
            start = text.find(needle)
            if start != -1:
                findings.append(
                    type(
                        "Tmp",
                        (),
                        {
                            "start": start,
                            "end": start + len(needle),
                            "text": needle,
                            "confidence": 0.9,
                            "label": label,
                            "reason": "sensitive block",
                        },
                    )()
                )
        return findings

    engine = PersonaEngine(
        backend=MockLLMBackend(response_factory=response_factory),
        backend_config=LLMBackendConfig(backend="mock"),
    )

    anonymize_result = engine.anonymize_document(
        input_path=input_path,
        password="secret-pass",
        out_dir=tmp_path,
        review=False,
    )

    censored = Document(anonymize_result.output_file)
    assert "[[P2|" in censored.paragraphs[0].text
    assert len(anonymize_result.entries) == 2

    restore_result = engine.restore_document(
        censored_path=Path(anonymize_result.output_file),
        map_path=Path(anonymize_result.map_file),
        password="secret-pass",
        out_dir=tmp_path,
    )

    restored = Document(restore_result.output_file)
    assert restored.paragraphs[0].text == "Alice Example works at Secret Corp."


def test_local_app_lists_documents_and_supports_review_and_preview_switch(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setenv("PERSONA_KEYSTORE_PATH", str(tmp_path / "keystore.json"))
    store = LocalDocumentStore(tmp_path / "workspace")
    engine = PersonaEngine(
        backend=MockLLMBackend(
            response_factory=lambda text: [
                type(
                    "Tmp",
                    (),
                    {
                        "start": text.find("Alice Example"),
                        "end": text.find("Alice Example") + len("Alice Example"),
                        "text": "Alice Example",
                        "confidence": 0.94,
                        "label": "person",
                        "reason": "personal identity block",
                    },
                )()
            ]
        ),
        backend_config=LLMBackendConfig(backend="mock"),
    )
    client = TestClient(create_web_app(engine=engine, store=store))

    source_path = tmp_path / "upload.docx"
    _build_docx(source_path, "Alice Example shares data.")

    with source_path.open("rb") as handle:
        response = client.post(
            "/documents/upload",
            files={"document": ("upload.docx", handle.read(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
            data={"password": "secret-pass"},
            follow_redirects=False,
        )

    assert response.status_code == 303
    detail_url = response.headers["location"]

    detail = client.get(detail_url)
    assert "Alice Example" in detail.text
    assert "Anteprima originale" in detail.text

    document_id = detail_url.rsplit("/", 1)[-1]
    record = store.load_document(document_id)
    first_match = record.hydrated_matches()[0]
    review_response = client.post(
        f"/documents/{document_id}/review",
        data={
            f"approve_{first_match.match_id}": "1",
            f"masked_{first_match.match_id}": "Masked Person",
        },
        follow_redirects=False,
    )
    assert review_response.status_code == 303

    anonymize_response = client.post(
        f"/documents/{document_id}/anonymize",
        data={"password": "secret-pass"},
        follow_redirects=False,
    )
    assert anonymize_response.status_code == 303

    anonymized_page = client.get(f"/documents/{document_id}?view=anonymized")
    assert "Anteprima anonimizzata" in anonymized_page.text
    assert "[[P2|" in anonymized_page.text
    assert "Review blocchi sensibili" in anonymized_page.text


def test_local_document_store_lists_saved_documents(tmp_path: Path) -> None:
    store = LocalDocumentStore(tmp_path / "workspace")
    record = StoredDocument(
        document_id="doc123",
        original_name="example.pdf",
        file_format=".pdf",
        created_at="2026-03-10T12:00:00+00:00",
        status="uploaded",
        original_path=str(tmp_path / "workspace" / "originals" / "doc123.pdf"),
    )
    store.save_document(record)

    listed = store.list_documents()

    assert len(listed) == 1
    assert listed[0].document_id == "doc123"
