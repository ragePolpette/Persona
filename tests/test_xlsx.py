from __future__ import annotations

import re
import zipfile
from pathlib import Path

import pytest
from openpyxl import load_workbook
from typer.testing import CliRunner

from persona import cli
from persona.documents import XlsxDocument, open_document
from persona.engine import analyze, apply, verify
from persona.exceptions import InputError
from persona.restore import placeholders_in, restore_segments
from persona.vault import Argon2Params, Vault
from tests.xlsx_fixtures import SENSITIVE, all_xml_text, build_xlsx


@pytest.fixture
def source(tmp_path: Path) -> Path:
    return build_xlsx(tmp_path / "clienti.xlsx")


def strip(placeholder: str) -> str:
    return placeholder.replace("[", "").replace("]", "")


def anonymize(source: Path, vault: Vault, out: Path) -> Path:
    document = open_document(source)
    result = apply(analyze(document.segments, vault), vault)
    document.write(out, result.edits)
    return out


def test_segments_reach_every_hiding_place(source: Path) -> None:
    ids = {s.id: s.text for s in open_document(source).segments}
    joined = "\n".join(ids.values())
    assert "Mario Rossi" in [t for i, t in ids.items() if "sharedStrings" in i]  # rich text over two runs
    assert "Contratto Giulia Marchetti" in ids.values()
    assert "ジュリア" in ids.values()  # phonetic guide
    assert "Acme S.r.l." in [t for i, t in ids.items() if i.endswith("#A2")]  # cell
    assert "Acme S.r.l." in [t for i, t in ids.items() if "workbook.xml#sheet" in i]  # sheet name
    assert any(i.endswith(":f") and "'Acme S.r.l.'!E2*2" in t for i, t in ids.items())  # formula
    assert any("comments" in i and "Chiamare Mario Rossi" in t for i, t in ids.items())
    assert any("rels" in i and "mailto:mario.rossi@example.test" in t for i, t in ids.items())  # hyperlink target
    assert "Riservato - Mario Rossi" in joined  # sheet header
    assert "Scegli Acme S.r.l." in joined  # validation prompt
    assert '"Acme S.r.l.,Beta S.p.A."' in joined  # validation list
    assert any("workbook.xml#name" in i and "Acme S.r.l." in t for i, t in ids.items())  # defined name
    assert "Referente Giulia Marchetti" in joined  # inline string
    assert "Sentire Mario Rossi" in joined  # threaded comment
    assert "Fatture Acme S.r.l." in joined  # title property


def test_nothing_sensitive_survives_anywhere_in_the_package(source: Path, vault: Vault, tmp_path: Path) -> None:
    before = all_xml_text(source)
    assert all(item in before for item in SENSITIVE)
    out = anonymize(source, vault, tmp_path / "out.xlsx")
    package = all_xml_text(out)
    for item in SENSITIVE:
        assert item not in package, f"leaked: {item!r}"
    for fragment in ("Mario", "Rossi", "Marchetti", "Acme", "Beta S.p.A.", "ジュリア"):
        assert fragment not in package, f"leaked fragment: {fragment!r}"
    assert verify({s.id: s.text for s in open_document(out).segments}, vault) == []


def test_sheet_names_and_formulas_stay_consistent_and_valid(source: Path, vault: Vault, tmp_path: Path) -> None:
    out = anonymize(source, vault, tmp_path / "out.xlsx")
    workbook = load_workbook(str(out))  # still a readable workbook
    first, second = workbook.sheetnames
    renamed = strip(vault.placeholder_for("AZIENDA", "Acme S.r.l."))
    assert first == renamed and re.fullmatch(r"AZIENDA_\d+", first)
    assert second == "Riepilogo"
    assert workbook["Riepilogo"]["B1"].value == f"='{renamed}'!E2*2"  # still points at the renamed sheet
    assert workbook.defined_names["Cliente1"].attr_text == f"'{renamed}'!$A$2"


def test_values_types_and_layout_are_preserved(source: Path, vault: Vault, tmp_path: Path) -> None:
    out = anonymize(source, vault, tmp_path / "out.xlsx")
    sheet = load_workbook(str(out)).worksheets[0]
    assert sheet["E2"].value == 12500 and sheet["F2"].value == 3384567890  # numbers stay numbers
    assert sheet["A2"].value.startswith("[AZIENDA_")
    assert sheet.column_dimensions["A"].width == 33
    assert load_workbook(str(out))["Riepilogo"]["B2"].value == "30 giorni"


def test_restore_gives_back_the_original(source: Path, vault: Vault, tmp_path: Path) -> None:
    out = anonymize(source, vault, tmp_path / "out.xlsx")
    censored = open_document(out)
    expected = set().union(*(placeholders_in(s.text) for s in censored.segments))
    edits, report = restore_segments(censored.segments, vault, expected=expected)
    restored = tmp_path / "restored.xlsx"
    censored.write(restored, edits)
    assert report.clean
    # phonetic guides of masked strings are dropped on purpose and not restored
    keep = lambda s: "docProps" not in s.id and not s.id.endswith(":ph")  # noqa: E731
    original = {s.id: s.text for s in open_document(source).segments if keep(s)}
    back = {s.id: s.text for s in open_document(restored).segments if keep(s)}
    assert back == original
    assert load_workbook(str(restored)).sheetnames == ["Acme S.r.l.", "Riepilogo"]


def test_identity_traces_are_scrubbed(source: Path, vault: Vault, tmp_path: Path) -> None:
    out = anonymize(source, vault, tmp_path / "out.xlsx")
    with zipfile.ZipFile(out) as archive:
        assert "docProps/thumbnail.jpeg" not in archive.namelist()
        assert b"thumbnail" not in archive.read("_rels/.rels")
        assert b"<author>Author</author>" in archive.read("xl/comments/comment1.xml")
        person = archive.read("xl/persons/person.xml")
        assert b'displayName="Author"' in person and b"mario" not in person
    assert load_workbook(str(out)).properties.creator in (None, "")


def test_numeric_identifiers_and_unreadable_content_are_flagged(tmp_path: Path) -> None:
    plain = open_document(build_xlsx(tmp_path / "a.xlsx"))
    assert any("numeric cell" in w and "Acme S.r.l.!F2" in w for w in plain.warnings)
    rich = open_document(build_xlsx(tmp_path / "b.xlsx", with_chart_and_drawing=True))
    assert any("chart" in w and "NOT anonymized" in w for w in rich.warnings)
    assert any("drawings" in w for w in rich.warnings)


def test_invalid_files_are_clean_errors(tmp_path: Path) -> None:
    bad = tmp_path / "bad.xlsx"
    bad.write_bytes(b"not a zip")
    with pytest.raises(InputError, match="not a valid XLSX"):
        open_document(bad)
    empty = tmp_path / "empty.xlsx"
    with zipfile.ZipFile(empty, "w") as archive:
        archive.writestr("hello.txt", "x")
    with pytest.raises(InputError, match="no xl/workbook.xml"):
        XlsxDocument(empty)


def test_ai_generated_workbooks_with_inline_strings_and_mangled_placeholders(vault: Vault, tmp_path: Path) -> None:
    vault.placeholder_for("AZIENDA", "Acme S.r.l.")
    vault.placeholder_for("PERSONA", "Mario Rossi")
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "AZIENDA_1"
    ws["A1"] = "Cliente: [azienda 1]"
    ws["A2"] = r"\[PERSONA\_1\]"
    ws["A3"] = "=\"x\"&'AZIENDA_1'!A1"
    path = tmp_path / "ai.xlsx"
    wb.save(str(path))
    document = open_document(path)
    edits, report = restore_segments(document.segments, vault)
    out = tmp_path / "ai.restored.xlsx"
    document.write(out, edits)
    restored = load_workbook(str(out))
    assert restored.sheetnames == ["Acme S.r.l."]
    sheet = restored["Acme S.r.l."]
    assert sheet["A1"].value == "Cliente: Acme S.r.l."
    assert sheet["A2"].value == "Mario Rossi"
    assert sheet["A3"].value == "=\"x\"&'Acme S.r.l.'!A1"
    assert report.clean


def test_xlsx_through_the_cli(source: Path, tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setenv("PERSONA_HOME", str(tmp_path / "home"))
    monkeypatch.setenv("PERSONA_PASSWORD", "pw")
    monkeypatch.setattr(Argon2Params.__init__, "__defaults__", (1, 8, 1, 32))
    runner = CliRunner()
    assert runner.invoke(cli.app, ["init", "-p", "x"]).exit_code == 0
    done = runner.invoke(cli.app, ["anonymize", str(source), "-p", "x"])
    assert done.exit_code == 0, done.output
    assert "numeric cell" in done.output  # the warning reaches the user
    anon = tmp_path / "clienti.anon.xlsx"
    assert "Acme" not in all_xml_text(anon)
    assert runner.invoke(cli.app, ["verify", str(anon), "-p", "x"]).exit_code == 0
    back = runner.invoke(cli.app, ["restore", str(anon), "--sent", str(anon), "-p", "x"])
    assert back.exit_code == 0, back.output
    assert load_workbook(str(tmp_path / "clienti.anon.restored.xlsx")).sheetnames == ["Acme S.r.l.", "Riepilogo"]


def test_rich_text_keeps_the_formatting_of_its_first_run(source: Path, vault: Vault, tmp_path: Path) -> None:
    out = anonymize(source, vault, tmp_path / "out.xlsx")
    shared = zipfile.ZipFile(out).read("xl/sharedStrings.xml").decode("utf-8")
    first = re.search(r"<si>(.*?)</si>", shared, re.S).group(1)
    assert "<b/>" in first and "[PERSONA_" in first and "Rossi" not in first
