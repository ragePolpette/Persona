from __future__ import annotations

from pathlib import Path

import pdfplumber
import pytest
from docx import Document
from openpyxl import Workbook, load_workbook
from reportlab.pdfgen import canvas

from persona.adapters.docx_adapter import DocxAdapter
from persona.adapters.pdf_adapter import PdfAdapter
from persona.adapters.xlsx_adapter import XlsxAdapter
from persona.exceptions import UnsupportedFormatError
from persona.models.entities import MapEntry, TextReplacement


def _make_replacement(segment_id: str, location: str, text: str, target: str, replacement_text: str) -> TextReplacement:
    start = text.index(target)
    return TextReplacement(
        start=start,
        end=start + len(target),
        replacement=replacement_text,
        token_id="TOKEN12345678",
        entity_type="PERSON",
        original_value=target,
        masked_value="Tario Xossi",
        location=location,
        segment_id=segment_id,
    )


def _make_entry(segment_id: str, location: str, original: str, placeholder: str, text: str) -> MapEntry:
    start = text.index(original)
    return MapEntry(
        token_id="TOKEN12345678",
        entity_type="PERSON",
        original_value=original,
        masked_value="Tario Xossi",
        placeholder=placeholder,
        location=location,
        segment_id=segment_id,
        start=start,
        end=start + len(original),
    )


def test_docx_adapter_replaces_paragraph_text(tmp_path: Path) -> None:
    input_path = tmp_path / "paragraph.docx"
    output_path = tmp_path / "paragraph.censored.docx"
    document = Document()
    document.add_paragraph("Contact Mario Rossi tomorrow.")
    document.save(input_path)

    adapter = DocxAdapter()
    segments = adapter.extract_segments(input_path)
    replacement = _make_replacement(segments[0].segment_id, segments[0].location, segments[0].text, "Mario Rossi", "[[P1|TOKEN12345678|Tario Xossi]]")

    adapter.apply_replacements(input_path, output_path, {segments[0].segment_id: [replacement]})

    censored = Document(output_path)
    assert "[[P1|TOKEN12345678|Tario Xossi]]" in censored.paragraphs[0].text


def test_docx_adapter_replaces_table_text(tmp_path: Path) -> None:
    input_path = tmp_path / "table.docx"
    output_path = tmp_path / "table.censored.docx"
    document = Document()
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Mario Rossi in table"
    document.save(input_path)

    adapter = DocxAdapter()
    segments = adapter.extract_segments(input_path)
    segment = next(item for item in segments if item.location.startswith("table0"))
    replacement = _make_replacement(segment.segment_id, segment.location, segment.text, "Mario Rossi", "[[P1|TOKEN12345678|Tario Xossi]]")

    adapter.apply_replacements(input_path, output_path, {segment.segment_id: [replacement]})

    censored = Document(output_path)
    assert "[[P1|TOKEN12345678|Tario Xossi]]" in censored.tables[0].cell(0, 0).text


def test_docx_adapter_handles_split_runs(tmp_path: Path) -> None:
    input_path = tmp_path / "split-runs.docx"
    output_path = tmp_path / "split-runs.censored.docx"
    restored_path = tmp_path / "split-runs.restored.docx"
    document = Document()
    paragraph = document.add_paragraph()
    paragraph.add_run("Ma")
    paragraph.add_run("rio ")
    paragraph.add_run("Ros")
    paragraph.add_run("si")
    document.save(input_path)

    adapter = DocxAdapter()
    segment = adapter.extract_segments(input_path)[0]
    placeholder = "[[P1|TOKEN12345678|Tario Xossi]]"
    replacement = _make_replacement(segment.segment_id, segment.location, segment.text, "Mario Rossi", placeholder)
    entry = _make_entry(segment.segment_id, segment.location, "Mario Rossi", placeholder, segment.text)

    adapter.apply_replacements(input_path, output_path, {segment.segment_id: [replacement]})
    stats = adapter.restore_file(output_path, restored_path, [entry])

    assert stats.restored_count == 1
    assert not stats.warnings
    restored = Document(restored_path)
    assert restored.paragraphs[0].text == "Mario Rossi"


def test_xlsx_adapter_updates_multiple_sheets_and_skips_formulas(tmp_path: Path) -> None:
    input_path = tmp_path / "workbook.xlsx"
    output_path = tmp_path / "workbook.censored.xlsx"
    workbook = Workbook()
    first = workbook.active
    first.title = "Sheet1"
    first["A1"] = "Mario Rossi"
    first["B1"] = "=A1"
    second = workbook.create_sheet("Sheet2")
    second["C3"] = "Contact Mario Rossi"
    workbook.save(input_path)

    adapter = XlsxAdapter()
    segments = adapter.extract_segments(input_path)
    replacements = {
        "Sheet1!A1": [_make_replacement("Sheet1!A1", "Sheet1!A1", "Mario Rossi", "Mario Rossi", "[[P1|TOKEN12345678|Tario Xossi]]")],
        "Sheet2!C3": [_make_replacement("Sheet2!C3", "Sheet2!C3", "Contact Mario Rossi", "Mario Rossi", "[[P1|TOKEN12345678|Tario Xossi]]")],
    }

    adapter.apply_replacements(input_path, output_path, replacements)

    censored = load_workbook(output_path)
    assert censored["Sheet1"]["A1"].value == "[[P1|TOKEN12345678|Tario Xossi]]"
    assert censored["Sheet1"]["B1"].value == "=A1"
    assert "[[P1|TOKEN12345678|Tario Xossi]]" in censored["Sheet2"]["C3"].value


def test_pdf_adapter_best_effort_round_trip(tmp_path: Path) -> None:
    input_path = tmp_path / "sample.pdf"
    output_path = tmp_path / "sample.censored.pdf"
    restored_path = tmp_path / "sample.restored.pdf"
    pdf_canvas = canvas.Canvas(str(input_path))
    pdf_canvas.drawString(72, 720, "Mario Rossi")
    pdf_canvas.save()

    adapter = PdfAdapter()
    segment = adapter.extract_segments(input_path)[0]
    placeholder = "[[P1|TOKEN12345678|Tario Xossi]]"
    replacement = _make_replacement(segment.segment_id, segment.location, segment.text, "Mario Rossi", placeholder)
    entry = _make_entry(segment.segment_id, segment.location, "Mario Rossi", placeholder, segment.text)

    adapter.apply_replacements(input_path, output_path, {segment.segment_id: [replacement]})
    stats = adapter.restore_file(output_path, restored_path, [entry])

    assert stats.restored_count == 1
    assert not stats.warnings
    with pdfplumber.open(str(restored_path)) as pdf:
        assert "Mario Rossi" in (pdf.pages[0].extract_text() or "")


def test_pdf_adapter_fails_cleanly_on_non_text_pdf(tmp_path: Path) -> None:
    input_path = tmp_path / "non-text.pdf"
    pdf_canvas = canvas.Canvas(str(input_path))
    pdf_canvas.rect(72, 700, 100, 100, fill=1)
    pdf_canvas.save()

    adapter = PdfAdapter()
    with pytest.raises(UnsupportedFormatError):
        adapter.extract_segments(input_path)
