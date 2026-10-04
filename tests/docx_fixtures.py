"""Builds DOCX files that hide sensitive data in all the places real Word files do."""

from __future__ import annotations

import io
import struct
import zipfile
import zlib
from pathlib import Path

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def _png_1px() -> bytes:
    def chunk(kind: bytes, data: bytes) -> bytes:
        body = kind + data
        return struct.pack(">I", len(data)) + body + struct.pack(">I", zlib.crc32(body) & 0xFFFFFFFF)

    header = struct.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 0)
    raw = b"\x00\xff\x00\x00"
    return b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", header) + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b"")


PNG_1PX = _png_1px()

SENSITIVE = [
    "Mario Rossi",
    "mario.rossi@example.test",
    "Tessitura Valdarno S.r.l.",
    "338 4567890",
    "Giulia Marchetti",
    "linkedin.com/in/mario-rossi",
]


def build_docx(path: Path, *, with_image: bool = False) -> Path:
    doc = Document()
    props = doc.core_properties
    props.author = "Mario Rossi"
    props.last_modified_by = "Giulia Marchetti"
    props.title = "Contratto Tessitura Valdarno S.r.l."

    section = doc.sections[0]
    section.header.paragraphs[0].text = "Riservato - Mario Rossi"
    section.footer.paragraphs[0].text = "Scrivere a mario.rossi@example.test"

    # a name split over runs with different formatting
    paragraph = doc.add_paragraph()
    paragraph.add_run("Contratto con ")
    paragraph.add_run("Mario ").bold = True
    paragraph.add_run("Rossi").bold = True
    paragraph.add_run(", telefono 338 4567890.")

    doc.add_paragraph("Fornitore: Tessitura Valdarno S.r.l., referente Dott.ssa Giulia Marchetti.")

    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Referente"
    table.rows[0].cells[1].text = "Mario Rossi"

    link_paragraph = doc.add_paragraph("Profilo: ")
    rel_id = doc.part.relate_to("https://www.linkedin.com/in/mario-rossi", RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "linkedin.com/in/mario-rossi"
    run.append(text)
    hyperlink.append(run)
    link_paragraph._p.append(hyperlink)

    doc.add_paragraph("Valori business: 12.500,00 euro entro 30 giorni.")
    if with_image:
        png = path.with_suffix(".png")
        png.write_bytes(PNG_1PX)
        doc.add_picture(str(png))
        png.unlink()
    doc.save(str(path))
    _inject(path)
    return path


def _inject(path: Path) -> None:
    textbox = (
        '<w:p xmlns:v="urn:schemas-microsoft-com:vml"><w:r><w:pict><v:shape><v:textbox><w:txbxContent>'
        "<w:p><w:r><w:t>Nel riquadro: Mario Rossi</w:t></w:r></w:p>"
        "</w:txbxContent></v:textbox></v:shape></w:pict></w:r></w:p>"
    )
    deleted = (
        '<w:p><w:del w:id="1" w:author="Mario Rossi" w:date="2026-01-01T00:00:00Z"><w:r>'
        "<w:delText>vecchio referente Mario Rossi</w:delText></w:r></w:del>"
        '<w:r><w:t xml:space="preserve"> testo rimasto</w:t></w:r></w:p>'
    )
    field = (
        '<w:p><w:r><w:fldChar w:fldCharType="begin"/></w:r><w:r><w:instrText xml:space="preserve">'
        ' HYPERLINK "mailto:mario.rossi@example.test" </w:instrText></w:r>'
        '<w:r><w:fldChar w:fldCharType="separate"/></w:r><w:r><w:t>scrivi qui</w:t></w:r>'
        '<w:r><w:fldChar w:fldCharType="end"/></w:r></w:p>'
    )
    footnotes = (
        f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:footnotes xmlns:w="{W_NS}">'
        '<w:footnote w:id="1"><w:p><w:r><w:t>Nota: contattare Giulia Marchetti</w:t></w:r></w:p></w:footnote>'
        "</w:footnotes>"
    )
    comments = (
        f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:comments xmlns:w="{W_NS}">'
        '<w:comment w:id="0" w:author="Mario Rossi" w:initials="MR">'
        "<w:p><w:r><w:t>Verificare con Mario Rossi</w:t></w:r></w:p></w:comment></w:comments>"
    )
    people = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<w15:people xmlns:w15="http://schemas.microsoft.com/office/word/2012/wordml">'
        '<w15:person w15:author="Mario Rossi"/></w15:people>'
    )

    with zipfile.ZipFile(path) as source:
        entries = {name: source.read(name) for name in source.namelist()}
        order = source.namelist()

    document = entries["word/document.xml"].decode("utf-8")
    document = document.replace("</w:body>", textbox + deleted + field + "</w:body>", 1)
    entries["word/document.xml"] = document.encode("utf-8")

    rels = entries["word/_rels/document.xml.rels"].decode("utf-8")
    rels = rels.replace(
        "</Relationships>",
        '<Relationship Id="rIdFn" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/footnotes" Target="footnotes.xml"/>'
        '<Relationship Id="rIdCm" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/comments" Target="comments.xml"/>'
        "</Relationships>",
    )
    entries["word/_rels/document.xml.rels"] = rels.encode("utf-8")
    entries["word/footnotes.xml"] = footnotes.encode("utf-8")
    entries["word/comments.xml"] = comments.encode("utf-8")
    entries["word/people.xml"] = people.encode("utf-8")
    entries["docProps/thumbnail.jpeg"] = b"\xff\xd8\xff thumbnail of the first page"

    root_rels = entries["_rels/.rels"].decode("utf-8")
    root_rels = root_rels.replace(
        "</Relationships>",
        '<Relationship Id="rIdTh" Type="http://schemas.openxmlformats.org/package/2006/relationships/metadata/thumbnail" Target="docProps/thumbnail.jpeg"/>'
        "</Relationships>",
    )
    entries["_rels/.rels"] = root_rels.encode("utf-8")

    content_types = entries["[Content_Types].xml"].decode("utf-8")
    additions = (
        '<Override PartName="/word/footnotes.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml"/>'
        '<Override PartName="/word/comments.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.comments+xml"/>'
    )
    if 'Extension="jpeg"' not in content_types:
        additions += '<Default Extension="jpeg" ContentType="image/jpeg"/>'
    entries["[Content_Types].xml"] = content_types.replace("</Types>", additions + "</Types>").encode("utf-8")

    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as target:
        for name in [*order, *(n for n in entries if n not in order)]:
            target.writestr(name, entries[name])
    path.write_bytes(buffer.getvalue())


def all_xml_text(path: Path) -> str:
    """Every part of the package as text, to prove nothing sensitive is left anywhere."""
    with zipfile.ZipFile(path) as archive:
        return "\n".join(
            archive.read(name).decode("utf-8", "ignore") for name in archive.namelist() if not name.endswith("/")
        )
