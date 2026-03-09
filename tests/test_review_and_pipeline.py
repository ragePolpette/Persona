from __future__ import annotations

import json
from pathlib import Path

import pytest
import spacy
from docx import Document
from rich.console import Console

from persona.core.pipeline import anonymize_file, restore_file
from persona.exceptions import ReviewAbortedError
from persona.models.entities import DetectionMatch
from persona.review.interactive import review_matches


def _build_test_model(tmp_path: Path) -> Path:
    model_path = tmp_path / "test_spacy_model"
    nlp = spacy.blank("en")
    ruler = nlp.add_pipe("entity_ruler")
    ruler.add_patterns([{"label": "PERSON", "pattern": "Mario Rossi"}])
    nlp.to_disk(model_path)
    return model_path


def test_review_flow_supports_edit_and_reject() -> None:
    matches = [
        DetectionMatch(
            match_id="m0001",
            segment_id="body/p0",
            location="body/p0",
            entity_type="PERSON",
            original_value="Mario Rossi",
            start=0,
            end=11,
            score=0.9,
            context="[Mario Rossi] works here",
            token_id="TOKEN12345678",
            masked_value="Tario Xossi",
            placeholder="[[P1|TOKEN12345678|Tario Xossi]]",
        ),
        DetectionMatch(
            match_id="m0002",
            segment_id="body/p0",
            location="body/p0",
            entity_type="EMAIL_ADDRESS",
            original_value="mario@example.com",
            start=20,
            end=37,
            score=0.9,
            context="mail [mario@example.com]",
            token_id="TOKEN87654321",
            masked_value="abcde@example.com",
            placeholder="[[P1|TOKEN87654321|abcde@example.com]]",
        ),
    ]
    answers = iter(["e", "Custom Mask", "r"])
    reviewed = review_matches(matches, prompt=lambda _: next(answers), console=Console(record=True))

    assert reviewed[0].approved is True
    assert reviewed[0].masked_value == "Custom Mask"
    assert reviewed[1].approved is False


def test_review_flow_can_abort() -> None:
    matches = [
        DetectionMatch(
            match_id="m0001",
            segment_id="body/p0",
            location="body/p0",
            entity_type="PERSON",
            original_value="Mario Rossi",
            start=0,
            end=11,
            score=0.9,
            context="[Mario Rossi] works here",
        )
    ]
    with pytest.raises(ReviewAbortedError):
        review_matches(matches, prompt=lambda _: "q", console=Console(record=True))


def test_anonymize_and_restore_docx_end_to_end(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    input_path = tmp_path / "source.docx"
    report_path = tmp_path / "anonymize-report.json"
    model_path = _build_test_model(tmp_path)
    monkeypatch.setenv("PERSONA_KEYSTORE_PATH", str(tmp_path / "keystore.json"))
    document = Document()
    document.add_paragraph("Mario Rossi email mario.rossi@example.com")
    document.save(input_path)

    anonymize_result = anonymize_file(
        input_path=input_path,
        password="top-secret",
        out_dir=tmp_path,
        review=False,
        enabled_entities=("PERSON", "EMAIL_ADDRESS"),
        report_json=report_path,
        spacy_model=str(model_path),
    )

    censored = Document(anonymize_result.output_file)
    assert "[[P2|" in censored.paragraphs[0].text
    assert Path(anonymize_result.map_file).exists()
    assert json.loads(report_path.read_text(encoding="utf-8"))["approved_matches"] == 2

    restore_result = restore_file(
        censored_path=Path(anonymize_result.output_file),
        map_path=Path(anonymize_result.map_file),
        password="top-secret",
        out_dir=tmp_path,
    )

    restored = Document(restore_result.output_file)
    assert restored.paragraphs[0].text == "Mario Rossi email mario.rossi@example.com"
    assert restore_result.restored_count == 2


def test_restore_strict_mode_leaves_tampered_placeholder(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    input_path = tmp_path / "source.docx"
    model_path = _build_test_model(tmp_path)
    monkeypatch.setenv("PERSONA_KEYSTORE_PATH", str(tmp_path / "keystore.json"))
    document = Document()
    document.add_paragraph("Mario Rossi")
    document.save(input_path)

    anonymize_result = anonymize_file(
        input_path=input_path,
        password="top-secret",
        out_dir=tmp_path,
        review=False,
        enabled_entities=("PERSON",),
        spacy_model=str(model_path),
    )

    censored_path = Path(anonymize_result.output_file)
    tampered = Document(censored_path)
    original_placeholder = anonymize_result.entries[0].placeholder
    tampered_placeholder = original_placeholder.replace("|", "!", 1)
    tampered.paragraphs[0].text = tampered.paragraphs[0].text.replace(original_placeholder, tampered_placeholder)
    tampered.save(censored_path)

    restore_result = restore_file(
        censored_path=censored_path,
        map_path=Path(anonymize_result.map_file),
        password="top-secret",
        out_dir=tmp_path,
    )

    restored = Document(restore_result.output_file)
    assert tampered_placeholder in restored.paragraphs[0].text
    assert restore_result.restored_count == 0
    assert restore_result.warnings
